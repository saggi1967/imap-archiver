"""Text-Extraktion aus Anhängen (PDF, DOCX, XLSX, Plaintext) für die Suche.

Jede Format-Funktion ist defensiv: schlägt die Extraktion fehl oder fehlt eine
optionale Bibliothek, wird der Anhang stillschweigend übersprungen (None), statt
die Indexierung abzubrechen.
"""

import io

_PLAIN_TYPES = {
    "text/plain",
    "text/csv",
    "text/markdown",
    "application/json",
    "application/xml",
    "text/xml",
}


def _ext(filename: str | None) -> str:
    if not filename or "." not in filename:
        return ""
    return filename.rsplit(".", 1)[1].lower()


def _kind(filename: str | None, content_type: str | None) -> str | None:
    ct = (content_type or "").lower()
    ext = _ext(filename)
    if ct == "application/pdf" or ext == "pdf":
        return "pdf"
    if "wordprocessingml" in ct or ext == "docx":
        return "docx"
    if "spreadsheetml" in ct or ext in ("xlsx", "xlsm"):
        return "xlsx"
    if ct in _PLAIN_TYPES or ct.startswith("text/") or ext in ("txt", "csv", "md", "json", "xml"):
        return "text"
    return None


def _from_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _from_docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    return "\n".join(p.text for p in document.paragraphs if p.text)


def _from_xlsx(data: bytes) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines: list[str] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                lines.append("\t".join(cells))
    wb.close()
    return "\n".join(lines)


def _from_text(data: bytes) -> str:
    return data.decode("utf-8", errors="replace")


_EXTRACTORS = {"pdf": _from_pdf, "docx": _from_docx, "xlsx": _from_xlsx, "text": _from_text}


def extract_text(
    filename: str | None, content_type: str | None, data: bytes, max_bytes: int
) -> str | None:
    """Liefert den Volltext eines Anhangs oder None (unsupported/zu groß/Fehler)."""
    if not data or len(data) > max_bytes:
        return None
    kind = _kind(filename, content_type)
    if kind is None:
        return None
    try:
        text = _EXTRACTORS[kind](data)
    except Exception:
        return None
    text = (text or "").strip()
    return text or None
